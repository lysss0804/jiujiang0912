"""Safe built-in conversion for DOCX documents; legacy formats need a controlled converter."""

from __future__ import annotations

from html import escape
from io import BytesIO

from docx import Document


def convert_docx(content: bytes, target_format: str) -> tuple[bytes, str, str]:
    if target_format not in {"txt", "markdown", "html"}:
        raise ValueError("built-in DOCX conversion supports txt, markdown or html")
    document = Document(BytesIO(content))
    blocks: list[str] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            blocks.append(text)
    for table in document.tables:
        for row in table.rows:
            blocks.append(" | ".join(cell.text.strip().replace("\n", " ") for cell in row.cells))
    if target_format == "txt":
        return "\n\n".join(blocks).encode("utf-8"), ".txt", "text/plain; charset=utf-8"
    if target_format == "markdown":
        return "\n\n".join(blocks).encode("utf-8"), ".md", "text/markdown; charset=utf-8"
    html = "\n".join(f"<p>{escape(block)}</p>" for block in blocks)
    return f"<!doctype html><html><body>{html}</body></html>".encode("utf-8"), ".html", "text/html; charset=utf-8"
